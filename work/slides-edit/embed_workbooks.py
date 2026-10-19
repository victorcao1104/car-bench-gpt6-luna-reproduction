import zipfile,xml.etree.ElementTree as E,posixpath as P,io,copy
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
C='http://schemas.openxmlformats.org/drawingml/2006/chart';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';PR='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types'
E.register_namespace('c',C);E.register_namespace('r',R);E.register_namespace('a','http://schemas.openxmlformats.org/drawingml/2006/main')
with zipfile.ZipFile('work/slides-edit/candidate.pptx') as z:d={n:z.read(n) for n in z.namelist()}
ct=E.fromstring(d['[Content_Types].xml'])
for n in list(d):
 if '/charts/' not in n or not n.endswith('.xml') or '/_rels/' in n:continue
 root=E.fromstring(d[n]);wb=Workbook();ws=wb.active;ws.title='Data';ws.cell(1,1,'Task')
 for j,ser in enumerate(root.findall('.//{'+C+'}ser'),2):
  col=get_column_letter(j);title=ser.find('{'+C+'}tx/{'+C+'}v').text;ws.cell(1,j,title)
  for tag,littag,reftag,cachetag,column in [('cat','strLit','strRef','strCache','A'),('val','numLit','numRef','numCache',col)]:
   holder=ser.find('{'+C+'}'+tag);lit=holder.find('{'+C+'}'+littag);pts=lit.findall('{'+C+'}pt')
   for pt in pts:
    v=pt.find('{'+C+'}v').text;ws.cell(int(pt.get('idx'))+2,1 if tag=='cat' else j,v if tag=='cat' else float(v))
   ref=E.Element('{'+C+'}'+reftag);E.SubElement(ref,'{'+C+'}f').text=f"'Data'!${column}$2:${column}${len(pts)+1}";cache=copy.deepcopy(lit);cache.tag='{'+C+'}'+cachetag;ref.append(cache);holder.remove(lit);holder.append(ref)
  tx=ser.find('{'+C+'}tx');tx.clear();ref=E.SubElement(tx,'{'+C+'}strRef');E.SubElement(ref,'{'+C+'}f').text=f"'Data'!${col}$1";cache=E.SubElement(ref,'{'+C+'}strCache');E.SubElement(cache,'{'+C+'}ptCount',val='1');pt=E.SubElement(cache,'{'+C+'}pt',idx='0');E.SubElement(pt,'{'+C+'}v').text=title
 out=io.BytesIO();wb.save(out)
 with zipfile.ZipFile(io.BytesIO(out.getvalue())) as xz:
  cells=E.fromstring(xz.read('xl/worksheets/sheet1.xml')); X='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; raw={cell.get('r'):cell.find('{'+X+'}v').text for cell in cells.iter('{'+X+'}c') if cell.find('{'+X+'}v') is not None}
 for j,ser in enumerate(root.findall('.//{'+C+'}ser'),2):
  for pt in ser.findall('{'+C+'}val/{'+C+'}numRef/{'+C+'}numCache/{'+C+'}pt'):pt.find('{'+C+'}v').text=raw[get_column_letter(j)+str(int(pt.get('idx'))+2)]
 dest=P.normpath(P.join(P.dirname(n),'../embeddings',P.basename(n)+'.xlsx'));d[dest]=out.getvalue()
 rel=E.Element('{'+PR+'}Relationships');E.SubElement(rel,'{'+PR+'}Relationship',Id='rIdWorkbook',Type=R+'/package',Target=P.relpath(dest,P.dirname(n)));rn=P.join(P.dirname(n),'_rels',P.basename(n)+'.rels');d[rn]=E.tostring(rel,encoding='utf-8',xml_declaration=True)
 ext=E.SubElement(root,'{'+C+'}externalData',{'{'+R+'}id':'rIdWorkbook'});E.SubElement(ext,'{'+C+'}autoUpdate',val='0');d[n]=E.tostring(root,encoding='utf-8',xml_declaration=True)
 E.SubElement(ct,'{'+CT+'}Override',PartName='/'+dest,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
E.register_namespace('',CT);d['[Content_Types].xml']=E.tostring(ct,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile('work/slides-edit/portable.pptx','w',zipfile.ZIP_DEFLATED) as z:
 for n,b in d.items():z.writestr(n,b)
print('Embedded two editable chart-data workbooks')
